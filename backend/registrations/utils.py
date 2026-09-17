"""
ITQAN — Registration Utilities
دوال مساعدة: توليد الرقم المرجعي + تصدير Excel
"""
import io
import random
from datetime import datetime
from pathlib import Path

from django.conf import settings
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def generate_reference_number():
    """
    توليد رقم مرجعي فريد لطلب التسجيل
    الصيغة: ITQ-{السنة}-{رقم عشوائي 4 أرقام}
    مثال: ITQ-2026-5821
    """
    from .models import Registration

    year = datetime.now().year
    while True:
        code = random.randint(1000, 9999)
        ref = f"ITQ-{year}-{code}"
        if not Registration.objects.filter(reference_number=ref).exists():
            return ref


def format_arabic_timestamp(dt=None):
    """
    تنسيق التاريخ والوقت بالعربي
    الخرج: "27 أغسطس 2026 - 11:35 م"
    """
    if dt is None:
        dt = datetime.now()

    # الأشهر بالعربي
    arabic_months = {
        1: 'يناير', 2: 'فبراير', 3: 'مارس', 4: 'أبريل',
        5: 'مايو', 6: 'يونيو', 7: 'يوليو', 8: 'أغسطس',
        9: 'سبتمبر', 10: 'أكتوبر', 11: 'نوفمبر', 12: 'ديسمبر'
    }

    month_name = arabic_months.get(dt.month, '')
    hour = dt.hour
    period = 'ص' if hour < 12 else 'م'
    if hour == 0:
        hour = 12
    elif hour > 12:
        hour -= 12

    return f"{dt.day} {month_name} {dt.year} - {hour:02d}:{dt.minute:02d} {period}"


def build_excel_workbook_from_queryset(queryset):
    """
    بناء ملف Excel كامل بتنسيق احترافي معتمد من منصة إتقان (يشمل الشعار، الترويسة، ألوان الهوية، والـ RTL)
    """
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'بيانات الطلاب المسجلين'
    ws.sheet_view.rightToLeft = True  # اتجاه الكتابة من اليمين لليسار (RTL)
    ws.sheet_view.showGridLines = True

    # 1. إدراج الشعار الرسمي
    logo_path = Path(settings.BASE_DIR).parent / 'assets' / 'images' / 'logo.png'
    if not logo_path.exists():
        logo_path = Path(settings.BASE_DIR) / 'assets' / 'images' / 'logo.png'
    if logo_path.exists():
        try:
            img = OpenpyxlImage(str(logo_path))
            img.width = 68
            img.height = 68
            ws.add_image(img, "A2")
        except Exception:
            pass

    # 2. الترويسة المؤسسية
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 18

    # اسم المنصة
    cell_title = ws["B2"]
    cell_title.value = "مـنـصـة إتـقـان للتعليم والتدريب التقني"
    cell_title.font = Font(name="Arial", size=15, bold=True, color="1E3A8A")
    cell_title.alignment = Alignment(horizontal="right", vertical="center")

    # العنوان الفرعي
    cell_sub = ws["B3"]
    cell_sub.value = "سجل طلبات التسجيل وبيانات الطلاب المسجلين عبر البوابة الإلكترونية"
    cell_sub.font = Font(name="Arial", size=11, bold=True, color="475569")
    cell_sub.alignment = Alignment(horizontal="right", vertical="center")

    # بيانات التاريخ والتصدير
    cell_meta = ws["B4"]
    now_str = format_arabic_timestamp()
    total_count = queryset.count() if hasattr(queryset, 'count') else len(queryset)
    cell_meta.value = f"تاريخ الاستخراج: {now_str}  |  إجمالي الطلبات: {total_count} طالب  |  نظام إتقان المعتمد"
    cell_meta.font = Font(name="Arial", size=9, italic=True, color="64748B")
    cell_meta.alignment = Alignment(horizontal="right", vertical="center")

    # 3. عناوين الأعمدة
    headers = [
        '#',
        'الرقم المرجعي',
        'اسم الطالب (عربي)',
        'اسم الطالب (إنجليزي)',
        'رقم الجوال / الواتساب',
        'معرف الدورة',
        'اسم البرنامج / الدورة التدريبية',
        'تاريخ الميلاد',
        'مكان الميلاد',
        'حالة الطلب',
        'تاريخ ووقت التقديم',
        'ملف السند المرفوع',
    ]

    header_row_idx = 6
    ws.row_dimensions[header_row_idx].height = 28

    header_font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')  # Slate Navy فخم
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    status_map = {
        'pending': 'قيد المراجعة ⏳',
        'approved': 'مقبول ✅',
        'rejected': 'مرفوض ❌',
    }

    # ألوان الشارات والحالات
    status_fills = {
        'approved': PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid'),
        'pending':  PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid'),
        'rejected': PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid'),
    }
    status_fonts = {
        'approved': Font(name='Arial', size=10, bold=True, color='15803D'),
        'pending':  Font(name='Arial', size=10, bold=True, color='B45309'),
        'rejected': Font(name='Arial', size=10, bold=True, color='B91C1C'),
    }

    data_font = Font(name='Arial', size=10)
    data_alignment_center = Alignment(horizontal='center', vertical='center')
    data_alignment_right = Alignment(horizontal='right', vertical='center')

    for idx, reg in enumerate(queryset, 1):
        birth_str = reg.birth_date.strftime('%Y-%m-%d') if reg.birth_date else '—'
        created_str = reg.created_at.strftime('%Y-%m-%d %H:%M') if reg.created_at else '—'
        file_name = reg.receipt_file.name if reg.receipt_file else 'لا يوجد ملف'

        row_data = [
            idx,
            reg.reference_number,
            reg.full_name_ar,
            reg.full_name_en or '—',
            reg.phone,
            reg.course_id,
            reg.course_title,
            birth_str,
            reg.birth_place or '—',
            status_map.get(reg.status, reg.status),
            created_str,
            file_name,
        ]

        row_num = header_row_idx + idx
        ws.row_dimensions[row_num].height = 22
        
        # تلوين متناوب للصفوف
        is_even = (idx % 2 == 0)
        row_fill = PatternFill(
            start_color='F8FAFC' if is_even else 'FFFFFF',
            end_color='F8FAFC' if is_even else 'FFFFFF',
            fill_type='solid'
        )

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=value)
            cell.font = data_font
            cell.alignment = data_alignment_right if col_idx in [3, 4, 7, 12] else data_alignment_center
            cell.border = thin_border
            cell.fill = row_fill

            # تمييز خلية الحالة بلون الشارة
            if col_idx == 10 and reg.status in status_fills:
                cell.fill = status_fills[reg.status]
                cell.font = status_fonts[reg.status]

    # ضبط العرض التلقائي للأعمدة
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for r in range(header_row_idx, header_row_idx + len(queryset) + 1):
            val = ws.cell(row=r, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ضبط عرض العمود الأول A لاستيعاب الشعار بشكل أنيق
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 14)

    return wb


def export_registration_to_excel(registration=None):
    """
    تحديث ملف Excel الإداري المحفوظ على السيرفر
    """
    from .models import Registration
    excel_path = Path(settings.EXPORTS_DIR) / 'registrations.xlsx'
    
    queryset = Registration.objects.all().order_by('-created_at')
    wb = build_excel_workbook_from_queryset(queryset)
    wb.save(str(excel_path))
    return str(excel_path)
